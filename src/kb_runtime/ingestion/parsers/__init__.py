"""
Parsers: bytes of a file in, markdown out (docs/ingestion.md, Parsers).

    parse(data, file_name) -> ParseResult(markdown, images, metadata)

A registry maps the extension to a parser. Parsers are synchronous and CPU-bound; the ingestion
pipeline runs them in a thread. A ParseError is deterministic: the item fails without retry.
Images found in a document are returned as (name, bytes) and referenced from the markdown as
`![name](images/name)`; the VLM subtask (phase 2) turns them into text.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass, field
from typing import Callable

MAX_PDF_RENDERED_PAGES = 200
MAX_SHEET_ROWS = 20_000
REJECTED = {".doc", ".xls", ".ppt", ".mp4", ".mov", ".avi", ".mkv", ".webm", ".mp3", ".wav", ".m4a", ".flac"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff"}


class ParseError(Exception):
    """The file cannot be parsed; retrying will not help."""


@dataclass
class ParseResult:
    """`metadata`: the document's own citation fields (title, authors, published, publisher, url),
    and `pages`: the offset where each page / slide starts in the final markdown (`parse()` turns
    the parsers' `page_anchors` into it, after normalising)."""

    markdown: str
    images: list[tuple[str, bytes]] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    page_anchors: list[str] = field(default_factory=list)       # the opening text of each page, in order


def page_offsets(markdown: str, anchors: list[str]) -> list[int]:
    """Where each page starts in `markdown`, found by its opening text (searched in order, so
    pages that start alike resolve to successive places). A page whose anchor is not found
    starts where the previous one did."""
    offsets, position = [], 0
    for anchor in anchors:
        key = normalise(anchor).strip()[:60]
        found = markdown.find(key, position) if key else -1
        if found >= 0:
            position = found
        offsets.append(position)
    return offsets


def _office_metadata(props) -> dict:
    """Citation fields from Office core properties (python-docx / python-pptx)."""
    out = {}
    if getattr(props, "title", None):
        out["title"] = props.title
    if getattr(props, "author", None):
        out["authors"] = [props.author]
    date = getattr(props, "created", None) or getattr(props, "modified", None)
    if date:
        out["published"] = date.date().isoformat()
    if getattr(props, "version", None):
        out["version"] = props.version
    return out


Parser = Callable[[bytes, str], ParseResult]
_REGISTRY: dict[str, Parser] = {}


def register(*extensions: str):
    def decorate(fn: Parser) -> Parser:
        for ext in extensions:
            _REGISTRY[ext] = fn
        return fn
    return decorate


def extension(file_name: str) -> str:
    return os.path.splitext(file_name.lower())[1]


def supported(file_name: str) -> bool:
    return extension(file_name) in _REGISTRY


def parse(data: bytes, file_name: str) -> ParseResult:
    ext = extension(file_name)
    if ext in REJECTED:
        raise ParseError(f"{ext} files are not supported (convert .doc/.xls/.ppt to the x formats; no audio or video)")
    parser = _REGISTRY.get(ext)
    if parser is None:
        raise ParseError(f"no parser for {ext or 'files without an extension'}")
    try:
        result = parser(data, file_name)
    except ParseError:
        raise
    except Exception as exc:     # a broken file, not a bug to retry
        raise ParseError(f"cannot parse {os.path.basename(file_name)}: {type(exc).__name__}: {exc}") from exc
    result.markdown = normalise(result.markdown)
    if result.page_anchors:
        result.metadata["pages"] = page_offsets(result.markdown, result.page_anchors)
    return result


def normalise(text: str) -> str:
    """CRLF to LF, control characters (but tab and newline) removed, trailing spaces trimmed."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(c for c in text if c in "\n\t" or ord(c) >= 32 or c == "\u3000")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip("\n") + "\n"


def decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def markdown_table(rows: list[list[str]]) -> str:
    rows = [[str(c if c is not None else "").replace("|", "\\|").replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    while width and all(len(r) < width or not r[width - 1] for r in rows):   # trailing empty columns
        width -= 1
    rows = [(r + [""] * width)[:width] for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# --- text formats -------------------------------------------------------------------------------
@register(".md", ".markdown", ".txt", ".text")
def _text(data: bytes, name: str) -> ParseResult:
    return ParseResult(decode(data))


@register(".csv")
def _csv(data: bytes, name: str) -> ParseResult:
    rows = list(csv.reader(io.StringIO(decode(data))))
    return ParseResult(markdown_table(rows))


@register(".json")
def _json(data: bytes, name: str) -> ParseResult:
    text = decode(data)
    try:
        text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except ValueError as exc:
        raise ParseError(f"invalid JSON: {exc}") from exc
    return ParseResult(f"```json\n{text}\n```")


def html_to_markdown(html: str) -> ParseResult:
    import trafilatura
    from markdownify import markdownify

    metadata = {}
    try:
        meta = trafilatura.extract_metadata(html)
        if meta is not None:
            if meta.title:
                metadata["title"] = meta.title
            if meta.author:
                metadata["authors"] = [a.strip() for a in meta.author.split(";") if a.strip()]
            if meta.date:
                metadata["published"] = meta.date
            if meta.sitename:
                metadata["publisher"] = meta.sitename
    except Exception:
        pass
    extracted = trafilatura.extract(html, output_format="markdown", include_tables=True, include_links=True,
                                    include_images=True, favor_recall=True)
    markdown = extracted or markdownify(html, heading_style="ATX")
    return ParseResult(markdown, metadata=metadata)


@register(".html", ".htm")
def _html(data: bytes, name: str) -> ParseResult:
    return html_to_markdown(decode(data))


# --- office and pdf -------------------------------------------------------------------------------
@register(".pdf")
def _pdf(data: bytes, name: str) -> ParseResult:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    parts, images, anchors, metadata = [], [], [], {}
    try:
        try:
            info = pdf.get_metadata_dict(skip_empty=True)
        except Exception:
            info = {}
        if info.get("Title"):
            metadata["title"] = info["Title"]
        if info.get("Author"):
            metadata["authors"] = [a.strip() for a in re.split(r";|\band\b", info["Author"]) if a.strip()]
        if info.get("CreationDate"):
            metadata["published"] = info["CreationDate"]
        for index in range(len(pdf)):
            page = pdf[index]
            text = page.get_textpage().get_text_range().strip()
            if text:
                parts.append(text)
                anchors.append(text)
            elif len(images) < MAX_PDF_RENDERED_PAGES:
                image_name = f"page-{index + 1}.png"
                buffer = io.BytesIO()
                page.render(scale=2).to_pil().save(buffer, format="PNG")
                images.append((image_name, buffer.getvalue()))
                parts.append(f"![page {index + 1}](images/{image_name})")
                anchors.append(parts[-1])
            else:
                anchors.append("")
    finally:
        pdf.close()
    if not parts:
        raise ParseError("the PDF has no pages")
    return ParseResult("\n\n".join(parts), images, metadata, anchors)


@register(".docx")
def _docx(data: bytes, name: str) -> ParseResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    blocks, images = [], []
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            image_name = os.path.basename(rel.target_ref)
            images.append((image_name, rel.target_part.blob))
            blocks.append(f"![{image_name}](images/{image_name})")
    body = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style.startswith("heading"):
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                body.append("#" * min(level, 6) + " " + text)
            elif style == "title":
                body.append("# " + text)
            elif "list" in style:
                body.append("- " + text)
            else:
                body.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            body.append(markdown_table([[cell.text for cell in row.cells] for row in table.rows]))
    return ParseResult("\n\n".join(body + blocks), images, _office_metadata(document.core_properties))


@register(".xlsx", ".xlsm")
def _xlsx(data: bytes, name: str) -> ParseResult:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if v is None else str(v) for v in row])
                if len(rows) >= MAX_SHEET_ROWS:
                    break
            table = markdown_table(rows)
            if table:
                parts.append(f"## {sheet.title}\n\n{table}")
    finally:
        book.close()
    return ParseResult("\n\n".join(parts))


@register(".pptx")
def _pptx(data: bytes, name: str) -> ParseResult:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    parts, images = [], []
    for number, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines = [f"## Slide {number}" + (f": {title}" if title else "")]
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append("  " * para.level + "- " + text)
            if getattr(shape, "has_table", False) and shape.has_table:
                lines.append(markdown_table([[c.text for c in row.cells] for row in shape.table.rows]))
            if shape.shape_type == 13 and hasattr(shape, "image"):     # picture
                image_name = f"slide{number}-{len(images) + 1}.{shape.image.ext}"
                images.append((image_name, shape.image.blob))
                lines.append(f"![{image_name}](images/{image_name})")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("> Notes: " + notes.replace("\n", " "))
        parts.append("\n".join(lines))
    return ParseResult("\n\n".join(parts), images, _office_metadata(deck.core_properties),
                       [p.split("\n", 1)[0] for p in parts])


@register(".epub")
def _epub(data: bytes, name: str) -> ParseResult:
    import tempfile

    import ebooklib
    from ebooklib import epub
    from markdownify import markdownify

    with tempfile.NamedTemporaryFile(suffix=".epub") as tmp:
        tmp.write(data)
        tmp.flush()
        book = epub.read_epub(tmp.name)
    parts = []
    for item_id, _ in book.spine:
        item = book.get_item_with_id(item_id)
        if item is not None and item.get_type() == ebooklib.ITEM_DOCUMENT:
            parts.append(markdownify(item.get_content().decode("utf-8", errors="replace"), heading_style="ATX"))
    metadata = {}
    for key, field_name in (("title", "title"), ("creator", "authors"), ("date", "published"),
                            ("publisher", "publisher")):
        values = [v[0] for v in book.get_metadata("DC", key) if v and v[0]]
        if values:
            metadata[field_name] = values if field_name == "authors" else values[0]
    return ParseResult("\n\n".join(p.strip() for p in parts if p.strip()), metadata=metadata)


@register(*IMAGE_EXTENSIONS)
def _image(data: bytes, name: str) -> ParseResult:
    image_name = os.path.basename(name)
    return ParseResult(f"![{image_name}](images/{image_name})", [(image_name, data)])

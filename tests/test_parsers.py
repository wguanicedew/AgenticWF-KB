"""Parsers (docs/ingestion.md, Parsers) and the SSRF-safe fetch."""

import asyncio
import http.server
import io
import threading

import pytest

from agenticwf.kb.runtime.ingestion.parsers import ParseError, normalise, parse


def minimal_pdf(text: str) -> bytes:
    """A one-page PDF with a text layer, written by hand (no PDF library needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def test_text_encodings():
    assert parse("héllo\r\nworld".encode("utf-8"), "a.txt").markdown == "héllo\nworld\n"
    assert parse("粒子".encode("gb18030"), "a.md").markdown == "粒子\n"
    assert parse("x".encode("utf-16"), "a.txt").markdown == "x\n"


def test_csv_and_json():
    md = parse(b"name,value\nalpha,1\nbeta,2\n", "t.csv").markdown
    assert md.splitlines()[:3] == ["| name | value |", "|---|---|", "| alpha | 1 |"]
    assert parse(b'{"a": [1, 2]}', "d.json").markdown.startswith("```json\n{")
    with pytest.raises(ParseError):
        parse(b"{broken", "d.json")


def test_html_drops_boilerplate():
    html = (b"<html><head><title>Pilot guide</title></head><body><nav>Home | About</nav>"
            b"<article><h1>Restarting the pilot</h1><p>Stop the service, then start it again with the "
            b"new configuration so that the stage-out retries are picked up.</p></article></body></html>")
    result = parse(html, "page.html")
    assert "Restarting the pilot" in result.markdown and "stage-out retries" in result.markdown


def test_pdf_text_layer():
    result = parse(minimal_pdf("Calorimeter calibration notes"), "doc.pdf")
    assert "Calorimeter calibration notes" in result.markdown


def test_docx_headings_lists_tables():
    import docx

    document = docx.Document()
    document.add_heading("Detector manual", level=1)
    document.add_paragraph("Introductory text.")
    document.add_paragraph("first item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "key", "value"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "gain", "1.5"
    buffer = io.BytesIO()
    document.save(buffer)
    md = parse(buffer.getvalue(), "m.docx").markdown
    assert "# Detector manual" in md and "- first item" in md and "| gain | 1.5 |" in md


def test_xlsx_sheets_as_tables():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Runs"
    sheet.append(["run", "events", None])
    sheet.append([1, 1000, None])
    book.create_sheet("Empty")
    buffer = io.BytesIO()
    book.save(buffer)
    md = parse(buffer.getvalue(), "r.xlsx").markdown
    assert "## Runs" in md and "| run | events |" in md and "Empty" not in md


def test_pptx_slides_and_notes():
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Status"
    slide.placeholders[1].text_frame.text = "All green"
    slide.notes_slide.notes_text_frame.text = "Mention the outage"
    buffer = io.BytesIO()
    deck.save(buffer)
    md = parse(buffer.getvalue(), "s.pptx").markdown
    assert "## Slide 1: Status" in md and "- All green" in md and "> Notes: Mention the outage" in md


def test_images_and_rejected_formats():
    result = parse(b"\x89PNG...", "photo.png")
    assert result.markdown.strip() == "![photo.png](images/photo.png)" and result.images[0][0] == "photo.png"
    for name in ("old.doc", "clip.mp4", "noext"):
        with pytest.raises(ParseError):
            parse(b"x", name)


def test_normalise_strips_control_characters():
    assert normalise("a\x00b\x07\r\nc  \n") == "ab\nc\n"


# --- SSRF -------------------------------------------------------------------------------------
class _Handler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/redirect-private":
            self.send_response(302)
            self.send_header("Location", "http://10.0.0.1/secret")
            self.end_headers()
            return
        body = b"<html><body><h1>Hello</h1><p>From the test server.</p></body></html>"
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


@pytest.fixture
def local_server():
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def test_private_addresses_are_refused(local_server, monkeypatch):
    from agenticwf.kb.runtime.security import ssrf
    from agenticwf.kb.runtime.security.ssrf import FetchError, fetch

    monkeypatch.setattr(ssrf, "whitelist", lambda: (set(), []))

    for url in (local_server + "/", "http://localhost/", "http://10.1.2.3/", "http://169.254.169.254/latest/",
                "ftp://example.org/", "file:///etc/passwd"):
        with pytest.raises(FetchError):
            asyncio.run(fetch(url))


def test_whitelisted_server_is_fetched_and_redirects_rechecked(local_server, monkeypatch):
    from agenticwf.kb.runtime.security import ssrf

    monkeypatch.setattr(ssrf, "whitelist", lambda: (set(), [__import__("ipaddress").ip_network("127.0.0.0/8")]))
    page = asyncio.run(ssrf.fetch(local_server + "/page"))
    assert page.status == 200 and page.content_type == "text/html" and b"Hello" in page.body
    with pytest.raises(ssrf.FetchError):          # the redirect goes to a private address
        asyncio.run(ssrf.fetch(local_server + "/redirect-private"))
    with pytest.raises(ssrf.FetchError):
        asyncio.run(ssrf.fetch(local_server + "/page", max_bytes=10))
